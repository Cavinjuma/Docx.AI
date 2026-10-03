from io import BytesIO

import pytest
from docx import Document as WordDocument

from app.ingestion import chunk_text, extract_text


def test_extract_text_decodes_utf8_text_and_markdown() -> None:
    plain_text = extract_text(b"\xef\xbb\xbfProject budget: 450000 dollars", ".txt")
    markdown = extract_text(b"# Policy\n\nVacation allowance is 21 days.", "md")

    assert plain_text == "Project budget: 450000 dollars"
    assert markdown == "# Policy\nVacation allowance is 21 days."


def test_extract_text_reads_docx_paragraphs_and_tables() -> None:
    document = WordDocument()
    document.add_paragraph("Northwind project budget")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Approved budget"
    table.cell(0, 1).text = "450000 dollars"
    output = BytesIO()
    document.save(output)

    extracted = extract_text(output.getvalue(), "docx")

    assert "Northwind project budget" in extracted
    assert "Approved budget | 450000 dollars" in extracted


def test_extract_text_rejects_documents_without_searchable_text() -> None:
    with pytest.raises(ValueError, match="No readable text"):
        extract_text(b"  \n", "txt")


def test_chunk_text_preserves_content_and_overlaps_chunks() -> None:
    chunks = chunk_text("one two three four five six", chunk_size=4, overlap=1)

    assert chunks == ["one two three four", "four five six"]

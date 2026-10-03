"""Document ingestion: text extraction, cleaning, and chunking."""

from __future__ import annotations

from io import BytesIO

from docx import Document as WordDocument
from pypdf import PdfReader


def extract_text(contents: bytes, file_type: str) -> str:
	"""Extract searchable text from a supported uploaded document."""
	normalized_type = file_type.lower().lstrip(".")
	if normalized_type not in {"txt", "md", "pdf", "docx"}:
		raise ValueError("Unsupported document type")

	try:
		if normalized_type in {"txt", "md"}:
			text = contents.decode("utf-8-sig")
		elif normalized_type == "pdf":
			reader = PdfReader(BytesIO(contents))
			text = "\n".join(page.extract_text() or "" for page in reader.pages)
		elif normalized_type == "docx":
			document = WordDocument(BytesIO(contents))
			paragraphs = [
				paragraph.text
				for paragraph in document.paragraphs
				if paragraph.text.strip()
			]
			table_rows = [
				" | ".join(cell.text for cell in row.cells)
				for table in document.tables
				for row in table.rows
			]
			text = "\n".join([*paragraphs, *table_rows])
	except Exception as error:
		message = "Unable to read this document. Ensure it is a valid, text-readable file."
		raise ValueError(message) from error

	text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
	if not text:
		raise ValueError("No readable text was found in the document.")
	return text


def chunk_text(text: str, chunk_size: int = 180, overlap: int = 30) -> list[str]:
	"""Split extracted document text into overlapping word-based chunks."""
	if chunk_size <= 0 or overlap < 0 or overlap >= chunk_size:
		raise ValueError("Chunk size must be positive and greater than the overlap")

	words = text.split()
	step = chunk_size - overlap
	return [" ".join(words[start : start + chunk_size]) for start in range(0, len(words), step)]


__all__ = ["chunk_text", "extract_text"]
